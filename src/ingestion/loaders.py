"""Document loaders for PDF, text, HTML, DOCX, and JSON sources."""

from __future__ import annotations

import json
from pathlib import Path

from langchain_community.document_loaders import (
    DirectoryLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document


def _load_docx(path: Path) -> list[Document]:
    try:
        import docx

        doc = docx.Document(str(path))
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        return [
            Document(
                page_content=text,
                metadata={"source": path.name, "file_path": str(path)},
            )
        ]
    except Exception as e:
        print(f"Error loading {path}: {e}")
        return []


def _load_json(path: Path) -> list[Document]:
    try:
        raw_text = path.read_text(encoding="utf-8", errors="ignore")
        data = json.loads(raw_text)
        formatted = json.dumps(data, indent=2)
        return [
            Document(
                page_content=formatted,
                metadata={"source": path.name, "file_path": str(path)},
            )
        ]
    except Exception as e:
        print(f"Error loading {path}: {e}")
        return []


def load_documents(raw_dir: Path) -> list[Document]:
    """Load all supported documents from the raw data directory."""
    if not raw_dir.exists():
        raise FileNotFoundError(f"Raw data directory not found: {raw_dir}")

    documents: list[Document] = []

    # PDFs
    pdf_files = list(raw_dir.rglob("*.pdf"))
    for path in pdf_files:
        try:
            loader = PyPDFLoader(str(path))
            docs = loader.load()
            for doc in docs:
                doc.metadata.setdefault("source", str(path.name))
                doc.metadata.setdefault("file_path", str(path))
            documents.extend(docs)
        except Exception as e:
            print(f"Error loading PDF {path}: {e}")

    # Text & Markdown
    for pattern in ("*.txt", "*.md"):
        for path in raw_dir.rglob(pattern):
            try:
                loader = TextLoader(str(path), encoding="utf-8")
                docs = loader.load()
                for doc in docs:
                    doc.metadata.setdefault("source", str(path.name))
                    doc.metadata.setdefault("file_path", str(path))
                documents.extend(docs)
            except Exception as e:
                print(f"Error loading text {path}: {e}")

    # HTML
    for pattern in ("*.html", "*.htm"):
        for path in raw_dir.rglob(pattern):
            try:
                text = path.read_text(encoding="utf-8", errors="ignore")
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": path.name, "file_path": str(path)},
                    )
                )
            except Exception as e:
                print(f"Error loading HTML {path}: {e}")

    # DOCX
    for path in raw_dir.rglob("*.docx"):
        documents.extend(_load_docx(path))

    # JSON
    for path in raw_dir.rglob("*.json"):
        documents.extend(_load_json(path))

    return documents


def load_directory_mixed(raw_dir: Path) -> list[Document]:
    """Optional convenience loader using DirectoryLoader for text files."""
    loader = DirectoryLoader(
        str(raw_dir),
        glob="**/*.txt",
        loader_cls=TextLoader,
        loader_kwargs={"encoding": "utf-8"},
        show_progress=False,
    )
    return loader.load()
